#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import load_workbook


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    text = text.replace("\u3000", " ")
    return re.sub(r"\s+", " ", text)


def normalize_code(value: Any) -> str:
    text = clean_text(value)
    if text.endswith(".0"):
        text = text[:-2]
    return text


def split_multi(value: Any) -> List[str]:
    text = clean_text(value)
    if not text:
        return []
    parts = re.split(r"[;；]\s*", text)
    return [part.strip() for part in parts if part and part.strip()]


def row_dict(headers: List[str], values: Iterable[Any]) -> Dict[str, Any]:
    data: Dict[str, Any] = {}
    for idx, value in enumerate(values):
        if idx >= len(headers):
            continue
        key = headers[idx]
        if key:
            data[key] = value
    return data


def parse_standard_sheet(path: Path) -> Dict[str, Dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["审查标准对照"]
    rules: Dict[str, Dict[str, Any]] = {}
    current_category = ""
    for values in ws.iter_rows(min_row=3, values_only=True):
        owner, category, code, content, score, method, note = values[:7]
        if category:
            current_category = clean_text(category)
        rule_code = normalize_code(code)
        if not rule_code:
            continue
        rules[rule_code] = {
            "rule_code": rule_code,
            "owner": clean_text(owner),
            "category": current_category,
            "rule_name": clean_text(content)[:80],
            "rule_content": clean_text(content),
            "score": score if isinstance(score, (int, float)) else clean_text(score),
            "review_method": clean_text(method),
            "review_note": clean_text(note),
            "review_points": [],
            "required_documents": [],
            "positive_examples": [],
            "negative_examples": [],
            "problem_patterns": [],
        }
    return rules


def parse_examples(path: Path, rules: Dict[str, Dict[str, Any]]) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["示例主表"]
    headers = [clean_text(cell) for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    for values in ws.iter_rows(min_row=2, values_only=True):
        raw = row_dict(headers, values)
        example_id = clean_text(raw.get("示例ID"))
        rule_code = normalize_code(raw.get("原评分点编号"))
        if not example_id or not rule_code:
            continue

        rule = rules.setdefault(
            rule_code,
            {
                "rule_code": rule_code,
                "owner": "",
                "category": "",
                "rule_name": clean_text(raw.get("原评分点内容"))[:80],
                "rule_content": clean_text(raw.get("原评分点内容")),
                "score": "",
                "review_method": "",
                "review_note": "",
                "review_points": [],
                "required_documents": [],
                "positive_examples": [],
                "negative_examples": [],
                "problem_patterns": [],
            },
        )

        review_point = clean_text(raw.get("审查要点"))
        if review_point and review_point not in rule["review_points"]:
            rule["review_points"].append(review_point)

        documents = split_multi(raw.get("关联文书名称"))
        for document in documents:
            if document not in rule["required_documents"]:
                rule["required_documents"].append(document)

        example = {
            "example_id": example_id,
            "rule_code": rule_code,
            "sub_rule_code": normalize_code(raw.get("子评分点编号")),
            "sub_rule_name": clean_text(raw.get("子评分点名称")),
            "review_point": review_point,
            "conclusion": clean_text(raw.get("示例结论")),
            "case_number": clean_text(raw.get("原始案号")),
            "case_short_name": clean_text(raw.get("案号简称")),
            "multi_document": clean_text(raw.get("是否涉及多文书")) == "是",
            "documents": documents,
            "pages": split_multi(raw.get("关联页码")),
            "image_refs": split_multi(raw.get("图片文件名或链接")),
            "key_excerpt": clean_text(raw.get("关键内容摘录")),
            "review_reason": clean_text(raw.get("示例说明")),
            "annotator": clean_text(raw.get("标注人")),
            "annotated_at": clean_text(raw.get("标注日期")),
            "remark": clean_text(raw.get("备注")),
        }

        if example["conclusion"] == "符合":
            rule["positive_examples"].append(example)
        elif example["conclusion"] == "不符合":
            rule["negative_examples"].append(example)


def parse_problem_workbook(path: Optional[Path], rules: Dict[str, Dict[str, Any]]) -> None:
    if not path or not path.exists():
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    if "问题梳理清单" not in wb.sheetnames:
        return
    ws = wb["问题梳理清单"]
    headers = [clean_text(cell) for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    for values in ws.iter_rows(min_row=2, values_only=True):
        raw = row_dict(headers, values)
        rule_code = normalize_code(raw.get("关联评分点（对照审查标准）"))
        problem = clean_text(raw.get("问题"))
        if not rule_code or not problem:
            continue
        rule = rules.setdefault(
            rule_code,
            {
                "rule_code": rule_code,
                "owner": "",
                "category": "",
                "rule_name": "",
                "rule_content": "",
                "score": "",
                "review_method": "",
                "review_note": "",
                "review_points": [],
                "required_documents": [],
                "positive_examples": [],
                "negative_examples": [],
                "problem_patterns": [],
            },
        )
        rule["problem_patterns"].append(
            {
                "problem": problem,
                "solution": clean_text(raw.get("对于该问题的个人思路（解决方案）")),
                "confirmed_response": clean_text(raw.get("回应记录（经袁老师确认）")),
                "resolved": clean_text(raw.get("是否解决")),
                "author": clean_text(raw.get("问题提出人")),
                "date": clean_text(raw.get("问题提出日期")),
            }
        )


def attach_stats(rules: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    stats = {
        "rule_count": len(rules),
        "review_point_count": 0,
        "required_document_count": 0,
        "positive_example_count": 0,
        "negative_example_count": 0,
        "problem_pattern_count": 0,
        "rules_with_negative_examples": 0,
    }
    document_counter: Dict[str, int] = defaultdict(int)
    for rule in rules.values():
        stats["review_point_count"] += len(rule.get("review_points") or [])
        stats["required_document_count"] += len(rule.get("required_documents") or [])
        stats["positive_example_count"] += len(rule.get("positive_examples") or [])
        stats["negative_example_count"] += len(rule.get("negative_examples") or [])
        stats["problem_pattern_count"] += len(rule.get("problem_patterns") or [])
        if rule.get("negative_examples"):
            stats["rules_with_negative_examples"] += 1
        for document in rule.get("required_documents") or []:
            document_counter[document] += 1
    stats["top_required_documents"] = [
        {"document": document, "count": count}
        for document, count in sorted(document_counter.items(), key=lambda item: (-item[1], item[0]))[:20]
    ]
    return stats


def sort_payload(payload: Dict[str, Any]) -> Dict[str, Any]:
    def sort_key(value: str) -> Any:
        return (0, int(value)) if value.isdigit() else (1, value)

    payload["rules"] = {code: payload["rules"][code] for code in sorted(payload["rules"], key=sort_key)}
    return payload


def build_pack(example_dir: Path, problem_workbook: Optional[Path]) -> Dict[str, Any]:
    workbook = example_dir / "000_案卷评查示例库搭建.xlsx"
    if not workbook.exists():
        raise FileNotFoundError(f"example workbook not found: {workbook}")

    rules = parse_standard_sheet(workbook)
    parse_examples(workbook, rules)
    parse_problem_workbook(problem_workbook, rules)

    payload = {
        "meta": {
            "source": "案卷评查示例库搭建",
            "example_workbook": workbook.name,
            "problem_workbook": problem_workbook.name if problem_workbook else "",
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "usage": "8011 后端在调用 LLM 前按规则编号检索并压缩为 llm_rule_knowledge。",
            "stats": attach_stats(rules),
        },
        "rules": rules,
    }
    return sort_payload(payload)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the LLM rule knowledge pack for the 8011 document review backend.")
    parser.add_argument("--example-dir", required=True, type=Path)
    parser.add_argument("--problem-workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("review/knowledge/rule_example_pack.json"))
    args = parser.parse_args()

    payload = build_pack(args.example_dir, args.problem_workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    stats = payload["meta"]["stats"]
    print(f"wrote {args.output}")
    print(
        "rules={rule_count}, positive={positive_example_count}, negative={negative_example_count}, "
        "problems={problem_pattern_count}".format(**stats)
    )


if __name__ == "__main__":
    main()
